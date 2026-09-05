"""Compile an audit-ready workpaper (.xlsx) for one reconciled invoice."""

from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import config

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
SEVERITY_FILL = {
    "critical": PatternFill("solid", fgColor="F4CCCC"),
    "high": PatternFill("solid", fgColor="FCE5CD"),
    "medium": PatternFill("solid", fgColor="FFF2CC"),
    "low": PatternFill("solid", fgColor="EAF1DD"),
}


def _write_header(sheet, headers, row=1):
    for col, name in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _autosize(sheet, max_width=60):
    for column in sheet.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width + 3, max_width)


def build(state):
    extracted = state.get("extracted") or {}
    po = state.get("po")
    vendor = state.get("vendor")
    bank_txn = state.get("bank_txn")
    exceptions = state.get("exceptions") or []

    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "AP Reconciliation Workpaper"
    summary["A1"].font = TITLE_FONT
    rows = [
        ("Prepared by", "AI-native AP agent (LangGraph)"),
        ("Prepared at (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("Run ID", state.get("run_id")),
        ("Source document", state.get("source_file")),
        ("Extraction method", extracted.get("extraction_method")),
        ("", ""),
        ("Vendor", extracted.get("vendor_name")),
        ("Vendor in master", "yes" if vendor else "NO"),
        ("Invoice number", extracted.get("invoice_number")),
        ("Invoice date", extracted.get("invoice_date")),
        ("Due date", extracted.get("due_date")),
        ("PO number", extracted.get("po_number") or "(none)"),
        ("Invoice total", extracted.get("total_amount")),
        ("Currency", extracted.get("currency")),
        ("", ""),
        ("Exceptions raised", len(exceptions)),
        ("Agent recommendation", (state.get("recommendation") or "").upper()),
        ("Disposition", "PENDING HUMAN APPROVAL"),
    ]
    for offset, (label, value) in enumerate(rows, start=3):
        summary.cell(row=offset, column=1, value=label).font = Font(bold=True)
        summary.cell(row=offset, column=2, value=value)

    narrative_row = len(rows) + 5
    summary.cell(row=narrative_row, column=1, value="Auditor narrative").font = Font(bold=True)
    cell = summary.cell(row=narrative_row + 1, column=1, value=state.get("narrative") or "")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary.merge_cells(start_row=narrative_row + 1, start_column=1, end_row=narrative_row + 8, end_column=6)
    _autosize(summary)

    match = wb.create_sheet("Three-Way Match")
    _write_header(match, ["Attribute", "Invoice", "Purchase order", "Bank feed", "Agrees?"])
    po_amount = po["amount"] if po else None
    bank_amount = abs(bank_txn["amount"]) if bank_txn else None
    invoice_total = extracted.get("total_amount")
    match_rows = [
        (
            "Counterparty",
            extracted.get("vendor_name"),
            (vendor or {}).get("name"),
            (bank_txn or {}).get("counterparty"),
            "yes" if vendor and extracted.get("vendor_name") else "NO",
        ),
        (
            "Reference",
            extracted.get("invoice_number"),
            (po or {}).get("po_number"),
            (bank_txn or {}).get("reference"),
            "yes" if po else "NO",
        ),
        (
            "Amount",
            invoice_total,
            po_amount,
            bank_amount,
            "yes" if po_amount is not None and invoice_total is not None
            and abs(invoice_total - po_amount) <= max(po_amount * config.AMOUNT_TOLERANCE_PCT, config.AMOUNT_TOLERANCE_ABS)
            else "NO",
        ),
        ("Date", extracted.get("invoice_date"), (po or {}).get("issued_date"), (bank_txn or {}).get("posted_date"), ""),
        ("Status", "received", (po or {}).get("status"), "settled" if bank_txn else "(no payment found)", ""),
    ]
    for row in match_rows:
        match.append(list(row))
    _autosize(match)

    exc_sheet = wb.create_sheet("Exceptions")
    _write_header(exc_sheet, ["#", "Code", "Severity", "Detail"])
    for idx, exc in enumerate(exceptions, start=1):
        exc_sheet.append([idx, exc["code"], exc["severity"], exc["detail"]])
        fill = SEVERITY_FILL.get(exc["severity"])
        if fill:
            for col in range(1, 5):
                exc_sheet.cell(row=idx + 1, column=col).fill = fill
    if not exceptions:
        exc_sheet.append([1, "none", "-", "No control exceptions identified."])
    _autosize(exc_sheet)

    items = wb.create_sheet("Line Items")
    _write_header(items, ["SKU", "Description", "Qty", "Unit price", "Amount"])
    for item in extracted.get("line_items") or []:
        items.append(
            [
                item.get("sku"),
                item.get("description"),
                item.get("qty"),
                item.get("unit_price"),
                item.get("amount"),
            ]
        )
    _autosize(items)

    source = wb.create_sheet("Source Document")
    source["A1"] = "Original document text (evidence)"
    source["A1"].font = Font(bold=True)
    for offset, line in enumerate((state.get("raw_text") or "").splitlines(), start=3):
        source.cell(row=offset, column=1, value=line)
    source.column_dimensions["A"].width = 90

    config.OUT_DIR.mkdir(parents=True, exist_ok=True)
    filename = f"workpaper_{extracted.get('invoice_number') or 'unknown'}_{state.get('run_id')}.xlsx"
    path = config.OUT_DIR / filename
    wb.save(path)
    return str(path)
